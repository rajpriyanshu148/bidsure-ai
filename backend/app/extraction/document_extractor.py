from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook


class DocumentExtractor:
    """Extract text from supported document formats."""

    @staticmethod
    def extract_pdf_pages(file_path: str) -> list[dict]:
        """Extract PDF text while preserving page numbers."""

        pages = []

        with fitz.open(file_path) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text().strip()

                pages.append(
                    {
                        "page_number": page_number,
                        "text": text,
                    }
                )

        return pages

    @staticmethod
    def extract_pdf(file_path: str) -> str:
        """Extract complete PDF text."""

        pages = DocumentExtractor.extract_pdf_pages(
            file_path
        )

        return "\n".join(
            page["text"]
            for page in pages
            if page["text"]
        ).strip()

    @staticmethod
    def extract_docx(file_path: str) -> str:
        """Extract text from a DOCX file."""

        document = Document(file_path)

        text = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text.append(
                    paragraph.text.strip()
                )

        for table in document.tables:
            for row in table.rows:
                row_text = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                text.append(
                    " | ".join(row_text)
                )

        return "\n".join(text).strip()

    @staticmethod
    def extract_xlsx(file_path: str) -> str:
        """Extract data from an XLSX workbook."""

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        text = []

        for sheet in workbook.worksheets:
            text.append(
                f"--- Sheet: {sheet.title} ---"
            )

            for row in sheet.iter_rows(
                values_only=True
            ):
                values = [
                    str(value).strip()
                    for value in row
                    if value is not None
                ]

                if values:
                    text.append(
                        " | ".join(values)
                    )

        workbook.close()

        return "\n".join(text).strip()

    @classmethod
    def extract(cls, file_path: str) -> str:
        """Automatically select the correct extractor."""

        extension = Path(
            file_path
        ).suffix.lower()

        if extension == ".pdf":
            return cls.extract_pdf(file_path)

        if extension == ".docx":
            return cls.extract_docx(file_path)

        if extension == ".xlsx":
            return cls.extract_xlsx(file_path)

        raise ValueError(
            f"Unsupported file format: {extension}"
        )